# -*- coding: utf-8 -*-
"""aloud-notes-ppt, step 1: copy a .pptx into the task folder, export it to PDF with PowerPoint, render every
slide to PNG and dump the slide text (titles, body, tables, speaker notes, picture count) to JSON and Markdown.

Only the COPY is opened in PowerPoint. The script closes only that copy and never calls Quit, so the user's
other open presentations stay open. The original deck is never opened or written.

Usage:  python extract_deck.py <deck.pptx> <task_folder>
Output: <task_folder>/temp/extract/{deck_copy.pptx, deck_copy.pdf, slides/slide_NN.png, slides.json, slides_text.md}
        <task_folder>/docs/notes_config.json   (template, only if it does not exist yet)
"""
import io, json, os, shutil, subprocess, sys

import pymupdf as fitz
from pptx import Presentation

sys.stdout = io.open(1, "w", encoding="utf-8", closefd=False)
if len(sys.argv) != 3:
    sys.exit("usage: python extract_deck.py <deck.pptx> <task_folder>")
SRC, ROOT = os.path.abspath(sys.argv[1]), os.path.abspath(sys.argv[2])
EXT = os.path.join(ROOT, "temp", "extract")
for d in ("sections", "figs", "docs", "archive", os.path.join("temp", "build"), os.path.join("temp", "extract", "slides")):
    os.makedirs(os.path.join(ROOT, d), exist_ok=True)
copy = os.path.join(EXT, "deck_copy.pptx")
shutil.copy2(SRC, copy)
print(f"📁 copied {SRC} -> {copy} (the original is never opened)")

pdf = os.path.join(EXT, "deck_copy.pdf")
ps = f"""
function Retry([scriptblock]$f) {{ for ($k = 0; $k -lt 30; $k++) {{ try {{ return & $f }} catch {{ if ($_.Exception.HResult -ne -2147418111) {{ throw }}; Start-Sleep -Milliseconds 1000 }} }}; throw "busy" }}
$pp = New-Object -ComObject PowerPoint.Application
$p = Retry {{ $pp.Presentations.Open('{copy}', $true, $false, $false) }}
Retry {{ $p.SaveAs('{pdf}', 32) }} | Out-Null
Retry {{ $p.Close() }} | Out-Null
"""
subprocess.run(["powershell", "-NoProfile", "-Command", ps], check=True)
print(f"🚀 exported {pdf}")

doc = fitz.open(pdf)
for i, pg in enumerate(doc, 1):
    pg.get_pixmap(dpi=110).save(os.path.join(EXT, "slides", f"slide_{i:02d}.png"))
print(f"✅ rendered {len(doc)} slide images")

prs = Presentation(copy); out, md = [], []
for i, sl in enumerate(prs.slides, 1):
    texts, tables, pics = [], [], 0
    for sh in sl.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            texts.append(sh.text_frame.text.strip())
        if getattr(sh, "has_table", False) and sh.has_table:
            tables.append([[c.text for c in r.cells] for r in sh.table.rows])
        if sh.shape_type == 13:
            pics += 1
    notes = sl.notes_slide.notes_text_frame.text.strip() if sl.has_notes_slide else ""
    words = sum(len(t.split()) for t in texts)
    out.append({"n": i, "texts": texts, "tables": tables, "pictures": pics, "notes": notes, "words": words})
    md += [f"## Slide {i}", ""] + [t.replace("\n", "  \n") + "\n" for t in texts]
    for t in tables:
        md += ["| " + " | ".join(r) + " |" for r in t] + [""]
    if notes:
        md += [f"> Speaker notes: {notes}", ""]
    md += [f"_pictures: {pics}_", ""]
json.dump(out, io.open(os.path.join(EXT, "slides.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
io.open(os.path.join(EXT, "slides_text.md"), "w", encoding="utf-8").write("\n".join(md))
print(f"✅ slides.json + slides_text.md written ({len(out)} slides)")

# skip-list candidates: the first slide, a closing "thank you" slide, and divider-like slides
# (at most two text boxes, no picture or table, a short heading and 8 to 30 words in all). Confirm on the images.
def divider_like(s):
    return (len(s["texts"]) <= 2 and not s["pictures"] and not s["tables"] and s["texts"]
            and len(s["texts"][0].split()) <= 8 and 8 <= s["words"] <= 30)
cand = [s["n"] for s in out if s["n"] == 1 or "thank you" in " ".join(s["texts"]).lower() or divider_like(s)]
print(f"🔍 skip candidates (title / divider / closing, check the images): {cand}")
cfg = os.path.join(ROOT, "docs", "notes_config.json")
if not os.path.exists(cfg):
    base = os.path.splitext(os.path.basename(SRC))[0]
    json.dump({"name": f"{base}_Read_Aloud_Notes", "title": f"{base}: Read-Aloud Notes",
               "subtitle": "One spoken paragraph per slide", "author": "", "affiliation": [], "year": "",
               "skip": cand, "parts": {}, "fontsize": 12, "words": [85, 140], "insight_words": [15, 45], "max_pages": 0,
               "slides_per_page": [0, 0]},
              io.open(cfg, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print(f"🔧 wrote config template {cfg}: fill title, author, skip, parts before building")
