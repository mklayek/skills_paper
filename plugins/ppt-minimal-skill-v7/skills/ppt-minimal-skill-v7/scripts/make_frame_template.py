# -*- coding: utf-8 -*-
"""Rebuild assets/frame_template.pptx from the user's class deck (run once, or when the class style changes).

    python make_frame_template.py <class_deck.pptx> <out_template.pptx>

Keeps only the FIRST slide (class title slide) and the LAST slide (Thank you slide) of a COPY of the
class deck, and turns the per-lecture texts into tokens that frame.py fills in:
    {{METHOD}}  big white title           ("IP Method")
    {{GROUP}}   red-on-yellow method group ("Electrical Methods")
    {{LECTURE}} two-digit lecture number   ("06")
Shapes get stable names so frame.py and qa.py can find them. The class deck itself is never written.
"""
import copy, re, sys
from pptx import Presentation
from pptx.oxml.ns import qn
from lxml import etree

SRC, OUT = sys.argv[1], sys.argv[2]
p = Presentation(SRC)
print(f"📁 class deck: {len(p.slides)} slides, keeping slide 1 and slide {len(p.slides)}")

# drop every slide except the first and the last (unreachable parts are not written on save)
sld_ids = p.slides._sldIdLst
for sid in list(sld_ids)[1:-1]:
    p.part.drop_rel(sid.get(qn('r:id'))); sld_ids.remove(sid)
title, thanks = p.slides[0], p.slides[1]


def merge_runs(para, text):
    """Put `text` in the first run of a paragraph (keeping its formatting) and drop the other runs."""
    runs = para._p.findall(qn('a:r'))
    runs[0].find(qn('a:t')).text = text
    for r in runs[1:]: para._p.remove(r)


def explicit_fill(sh, color, line=None):
    """Replace a theme-style fill (p:style fillRef) with an explicit color, so any theme renders it the same."""
    spPr = sh._element.spPr
    for tag in ('a:solidFill', 'a:noFill', 'a:ln'):
        for e in spPr.findall(qn(tag)): spPr.remove(e)
    geom = spPr.find(qn('a:prstGeom'))
    f = etree.SubElement(spPr, qn('a:solidFill')); etree.SubElement(f, qn('a:srgbClr')).set('val', color)
    geom.addnext(f)
    ln = etree.SubElement(spPr, qn('a:ln'))
    if line: ln.set('w', '38100'); s = etree.SubElement(ln, qn('a:solidFill')); etree.SubElement(s, qn('a:srgbClr')).set('val', line)
    else: etree.SubElement(ln, qn('a:noFill'))
    f.addnext(ln)
    st = sh._element.find(qn('p:style'))
    if st is not None: sh._element.remove(st)


names = {}
for sh in title.shapes:
    t = sh.text_frame.text if sh.has_text_frame else ''
    if sh.name == 'Rectangle 1': sh.name = 'Frame'
    elif t.startswith('Dr.'): sh.name = 'FrameAuthor'
    elif 'Lecture' in t: sh.name = 'FrameTitle'
    elif 'GSE' in t: sh.name = 'FrameCourse'; explicit_fill(sh, '70AD47', 'FFFFFF')
    elif sh.shape_type == 13:
        ar = sh.width / sh.height
        sh.name = 'TitleCorner' if sh.top < 914400 else ('TitleBook' if ar < 1 else 'TitlePhoto')
    names[sh.name] = sh
print("🔧 title slide shapes:", list(names))

# tokens in the title box: paragraph 1 = METHOD <br> GROUP |, paragraph 2 = ... Lecture NN
tb = names['FrameTitle'].text_frame
p1 = tb.paragraphs[0]._p
runs = p1.findall(qn('a:r'))
runs[0].find(qn('a:t')).text = '{{METHOD}}'
grp = [r for r in runs if r.find(qn('a:t')).text.strip() not in ('|',) and r is not runs[0]]
grp[0].find(qn('a:t')).text = '{{GROUP}}'
for r in grp[1:]: p1.remove(r)
runs2 = tb.paragraphs[1]._p.findall(qn('a:r'))
lec = [r for r in runs2 if 'Lecture' in (r.find(qn('a:t')).text or '')][0]
lec.find(qn('a:t')).text = re.sub(r'Lecture\s*\d*', 'Lecture {{LECTURE}}', lec.find(qn('a:t')).text)
for r in runs2[runs2.index(lec) + 1:]: tb.paragraphs[1]._p.remove(r)

for sh in thanks.shapes:
    t = sh.text_frame.text if sh.has_text_frame else ''
    if sh.name == 'Rectangle 1': sh.name = 'Frame'
    elif t.startswith('Thank'): sh.name = 'FrameThanks'
    elif t.startswith('Dr.'): sh.name = 'FrameAuthor'
    elif sh.shape_type == 13: sh.name = 'ThanksCollage'
print("🔧 thank-you slide shapes:", [s.name for s in thanks.shapes])
p.save(OUT)
print(f"✅ {OUT} written")
