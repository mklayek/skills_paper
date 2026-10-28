# -*- coding: utf-8 -*-
"""Checks the built deck against the house rules (v7: all arrow chains, per-topic titles, frame slides).

Run on Final.pptx (title slide + content + Thank you slide). Deck.pptx works too, the frame checks
then report that the frame slides are missing."""
import re, sys, zipfile
from pptx import Presentation
from pptx.util import Pt

FN = sys.argv[1] if len(sys.argv) > 1 else 'Deck.pptx'
# v7 symbol line: Times New Roman 18 pt, 'sym: meaning ; sym: meaning' (its ' ; ' separators are allowed)
# v7 continuation bullet: the part of a chain that would have wrapped, lvl="1" on the paragraph
is_cont = lambda pa: pa._p.pPr is not None and pa._p.pPr.get('lvl') == '1'
is_sym = lambda pa: bool(pa.runs) and pa.runs[0].font.size is not None and pa.runs[0].font.size.pt == 18 \
    and pa.runs[0].font.name == 'Times New Roman'
p = Presentation(FN)
fail = []
ok = lambda m: print('  OK   ' + m)
from pptx.oxml.ns import qn


def is_frame(sl):
    return any(sh.name == 'Frame' for sh in sl.shapes)


def is_divider(sl):
    bg = sl._element.find('.//' + qn('p:bg'))
    return bg is not None and bg.find('.//' + qn('a:blipFill')) is not None


FRAME = {i for i, sl in enumerate(p.slides, 1) if is_frame(sl)}
DIV = {i for i, sl in enumerate(p.slides, 1) if is_divider(sl)}


def bad(m):
    fail.append(m); print('  FAIL ' + m)


# ---------------------------------------------------------------- structure --
print('STRUCTURE')
boxes_bad = []
for i, sl in enumerate(p.slides, 1):
    if i in FRAME: continue
    names = [sh.name for sh in sl.shapes if sh.has_text_frame]
    tx = [sh for sh in sl.shapes if sh.has_text_frame]
    if names.count('Title') != 1: boxes_bad.append(f'slide {i}: {names.count("Title")} Title')
    if names.count('Body') != 1: boxes_bad.append(f'slide {i}: {names.count("Body")} Body')
    extra = [sh for sh in tx if sh.name not in ('Title', 'Body', 'TopicLabel', 'Caption')]
    # an image caption is the only other text box allowed
    for e in extra:
        t = e.text_frame.text.strip()
        if not t:          # the image border rectangle carries an empty text frame
            continue
        runs = [r for para in e.text_frame.paragraphs for r in para.runs]
        cap = bool(runs) and all(r.font.italic and r.font.size and r.font.size.pt == 12
                                 for r in runs)
        if not (cap or t.startswith('Section ')):
            boxes_bad.append(f'slide {i}: unexpected text box {t[:40]!r}')
print(f'  slides: {len(p.slides)}')
bad('; '.join(boxes_bad)) if boxes_bad else ok('every slide has exactly one Title and one Body, no stray boxes')

# ------------------------------------------------------------------ titles --
print('TITLES')
titles = []
for i, sl in enumerate(p.slides, 1):
    t = [sh for sh in sl.shapes if sh.name == 'Title']
    titles.append((i, t[0].text_frame.text.strip() if t else ''))
content = [(i, t) for i, t in titles if i not in FRAME and i not in DIV]
divider = [(i, t) for i, t in titles if i in DIV]
numbered = [(i, t) for i, t in content if re.match(r'^[◆■●]\s+\d', t)]
secnum = [(i, t) for i, t in content if re.search(r'\b\d+\.\d+\b', t)]
ok(f'{len(content)} content titles, {len(divider)} dividers, {len(FRAME)} frame slides')
bad(f'titles with a leading number: {numbered}') if numbered else ok('no numbered titles')
bad(f'titles with a source section number: {secnum}') if secnum else ok('no source section numbers in titles')

# ------------------------------------------- slide numbers and footers (none) --
print('NO SLIDE NUMBERS / FOOTERS')
with zipfile.ZipFile(FN) as z:
    sld = [n for n in z.namelist() if n.startswith('ppt/slides/slide')]
    fld = [n for n in sld if b'<a:fld' in z.read(n) and b'type="slidenum"' in z.read(n)]
bad(f'slide-number fields in {fld}') if fld else ok('no slide-number field anywhere')

# ------------------------------------------------- equation numbers (none) --
print('NO EQUATION NUMBERS')
eqno = []
for i, sl in enumerate(p.slides, 1):
    for sh in sl.shapes:
        if not sh.has_text_frame: continue
        for para in sh.text_frame.paragraphs:
            txt = para.text
            if re.search(r'\(\s*\d+[a-e]?\s*\)\s*$', txt) and not re.search(r'(Fig|Table)', txt):
                eqno.append((i, txt[-40:]))
bad(f'trailing equation numbers: {eqno[:4]}') if eqno else ok('no equation numbers on any line')
chap = [(i, pa.text) for i, sl in enumerate(p.slides, 1) for sh in sl.shapes
        if sh.has_text_frame for pa in sh.text_frame.paragraphs
        if re.search(r'\(\d+\.\d+\)', pa.text)]
bad(f'chapter-prefixed numbers: {chap[:3]}') if chap else ok('no chapter-prefixed equation numbers')

# -------------------------------------------------------------- raw markup --
print('NO RAW MARKUP')
raw = [(i, pa.text[:50]) for i, sl in enumerate(p.slides, 1) for sh in sl.shapes
       if sh.has_text_frame for pa in sh.text_frame.paragraphs
       if '_{' in pa.text or '^{' in pa.text or '**' in pa.text or re.search(r'[A-Za-z]_[A-Za-zρθφ]', pa.text)]
bad(f'raw markup left: {raw[:4]}') if raw else ok('no _{ ^{ or ** left in the text')

# --------------------------------------------------------------- US English --
print('US ENGLISH')
WHITE = set('noise source four hour contour precise rise expertise promise filled controlled '
            'called drilling drilled flour pour your our tour raised praised noised poised '
            'revised advised devised comprised supervised'.split())
pat = re.compile(r'\b\w+(?:isation|ised|ising|yse|ysed|olour|aviour|metre|entre|entres|'
                 r'ramme|ibre|itre)\b|\bgrey\b|\bwhilst\b|\bamongst\b|\balum inium\b', re.I)
hits = []
for i, sl in enumerate(p.slides, 1):
    for sh in sl.shapes:
        if sh.has_text_frame:
            for m in pat.finditer(sh.text_frame.text):
                if m.group(0).lower() not in WHITE: hits.append((i, m.group(0)))
        if sh.has_table:
            for r in sh.table.rows:
                for c in r.cells:
                    for m in pat.finditer(c.text):
                        if m.group(0).lower() not in WHITE: hits.append((i, m.group(0)))
bad(f'British spellings: {sorted(set(h[1] for h in hits))}') if hits else ok('US English throughout')

# --------------------------------------------------------------- subtitles --
print('SUBTITLES (v2: only where they add information)')
sub = lambda pa: bool(pa.runs) and bool(pa.runs[0].font.underline)
nsub = [i for i, sl in enumerate(p.slides, 1) for sh in sl.shapes
        if sh.name == 'Body' and sh.text_frame.paragraphs and sub(sh.text_frame.paragraphs[0])]
ok(f'{len(nsub)} slides keep a subtitle (title slide and named case histories)')
defsub = [i for i, sl in enumerate(p.slides, 1) for sh in sl.shapes if sh.name == 'Body'
          and sh.text_frame.paragraphs and sub(sh.text_frame.paragraphs[0])
          and sh.text_frame.paragraphs[0].text.lower().startswith('definition')]
bad(f'"Definition:" subtitles on slides {defsub}') if defsub else ok('no "Definition:" subtitles (v7)')
ncont = sum(1 for sl in p.slides for sh in sl.shapes if sh.name == 'Body' for pa in sh.text_frame.paragraphs if is_cont(pa))
ok(f'{ncont} continuation bullets (chains split so that no bullet wraps)')

# ------------------------------------------------------------ title fonts --
print('TITLES (v7: one font + dark color per topic, centered start, left continued, light or no fill)')


def lum(hx):
    r, g, b_ = (int(hx[k:k + 2], 16) / 255 for k in (0, 2, 4))
    return 0.2126 * r + 0.7152 * g + 0.0722 * b_


MARKS = '◆■●▶❖✦▲◉➤✿'


def strip(t):
    t = t.lstrip('★').strip()
    t = t[1:].strip() if t[:1] in MARKS else t
    return re.sub(r' \(continued\)$', '', t)


tbad, seen, prev, nfill, noval, nlab, fonts_order, marks_order, cur_mark = [], {}, None, 0, 0, 0, [], [], None
for i, sl in enumerate(p.slides, 1):
    if i in FRAME or i in DIV:
        prev = None if i in DIV else prev
        continue
    tt = [sh for sh in sl.shapes if sh.name == 'Title'][0]
    lab = [sh for sh in sl.shapes if sh.name == 'TopicLabel']
    oval = tt.shape_type == 1 and 'ellipse' in tt._element.xml[:3000]
    para = tt.text_frame.paragraphs[0]
    r = para.runs[-1]
    sig = (r.font.name, str(r.font.color.rgb))
    topic = lab[0].text_frame.text.strip() if lab else strip(tt.text_frame.text.strip())
    r0 = para.runs[0]
    mark = (r0.text.strip(), str(r0.font.color.rgb)) if len(para.runs) > 1 and r0.text.strip() in MARKS else None
    first = topic != prev
    if lab and first: tbad.append((i, 'black label on the first slide of a topic'))
    if lab:
        nlab += 1
        lr = lab[0].text_frame.paragraphs[0].runs[0]
        if str(lab[0].fill.fore_color.rgb) != '000000' or str(lr.font.color.rgb) != 'FFFFFF':
            tbad.append((i, 'topic label is not white text on black'))
    if oval:
        noval += 1
    else:
        al = para.alignment
        want = 2 if first else 1                     # PP_ALIGN.CENTER = 2, LEFT = 1
        if al is not None and int(al) != want:
            tbad.append((i, ('start of topic must be centered' if first else 'continued topic must be left-aligned'), topic[:30]))
        if tt.fill.type == 1:
            nfill += 1
            if lum(str(tt.fill.fore_color.rgb)) < 0.8: tbad.append((i, 'title fill is not light', str(tt.fill.fore_color.rgb)))
    if lum(sig[1]) > 0.35: tbad.append((i, 'title text color is not dark', sig[1]))
    if first:
        if fonts_order and fonts_order[-1] == sig[0]: tbad.append((i, 'same title font as the previous topic', sig[0]))
        if mark and marks_order and marks_order[-1] == mark: tbad.append((i, 'same title bullet as the previous topic', mark))
        fonts_order.append(sig[0]); seen[(i, topic)] = sig
        cur = sig
        cur_mark = mark                                # None on an oval topic start
        if mark: marks_order.append(mark)
    elif sig != cur:
        tbad.append((i, 'font/color differs within one topic', sig, cur))
    elif mark and cur_mark and mark != cur_mark:
        tbad.append((i, 'title bullet differs within one topic', mark, cur_mark))
    elif mark and not cur_mark:
        cur_mark = mark; marks_order.append(mark)      # topic opened with an oval: its bullet starts here
    prev = topic
bad(f'title rule breaks: {tbad[:5]}') if tbad else ok(f'{len(fonts_order)} topics, fonts follow the topic rule')
ok(f'title bullets per topic: {" ".join(m[0] for m in marks_order)}')
ok(f'{len(set(fonts_order))} distinct title fonts, {nfill} filled title boxes, {noval} oval titles, {nlab} black topic labels')
if len(fonts_order) >= 3 and len(set(fonts_order)) < 3: bad('fewer than 3 distinct title fonts')
if nfill > len(content) / 2: print('  NOTE more than half of the title boxes are filled: v7 keeps most of them unfilled')

# ------------------------------------------------- semicolons and em dashes --
print('WRITING (no semicolons, no em dashes)')
pun = []
for i, sl in enumerate(p.slides, 1):
    for sh in sl.shapes:
        texts = [pa.text for pa in sh.text_frame.paragraphs if not is_sym(pa)] if sh.has_text_frame else []
        if sh.has_table: texts += [c.text for r in sh.table.rows for c in r.cells]
        for x in texts:
            if ';' in x or '—' in x: pun.append((i, x[:60]))
bad(f'semicolons or em dashes: {pun[:4]}') if pun else ok('no semicolons and no em dashes in slide text')

# ------------------------------------------------------------ bullet fonts --
print('BULLET FONT CYCLE')
FONTS = ['Times New Roman', 'Comic Sans MS', 'Bookman Old Style']
wrong = []
for i, sl in enumerate(p.slides, 1):
    b = [sh for sh in sl.shapes if sh.name == 'Body']
    if not b: continue
    if i in FRAME or i in DIV: continue      # divider and frame slides carry no bullets
    k = 0
    for j, pa in enumerate(b[0].text_frame.paragraphs):
        if not pa.runs or (j == 0 and sub(pa)): continue
        f = pa.runs[0].font
        if f.name in ('Cambria Math', 'Cambria', 'Arial') or is_sym(pa) or is_cont(pa): continue
        if f.name != FONTS[k % 3]: wrong.append((i, j, f.name, FONTS[k % 3]))
        if f.size and f.size != Pt(22): wrong.append((i, j, 'size', f.size.pt))
        k += 1
bad(f'bullet font/size breaks: {wrong[:4]}') if wrong else ok('bullets follow the 3-font cycle at 22 pt')

# ------------------------------------------------------------ arrow chains --
print('ARROW CHAINS (v7: every bullet is a chain)')
chain_sl, prose_sl, mixdir, samecol, mixed = [], [], [], [], []
for i, sl in enumerate(p.slides, 1):
    if i in FRAME or i in DIV: continue
    b = [sh for sh in sl.shapes if sh.name == 'Body']
    if not b: continue
    nch = npl = 0
    paras = b[0].text_frame.paragraphs
    for j, pa in enumerate(paras):
        if not pa.runs or (j == 0 and sub(pa)): continue
        if pa.runs[0].font.name in ('Cambria Math', 'Cambria', 'Arial') or is_sym(pa) or is_cont(pa): continue
        # a chain = this paragraph plus the continuation bullets that follow it (v7 one-line bullets)
        grp = [pa]
        for q in paras[j + 1:]:
            if not is_cont(q): break
            grp.append(q)
        t = ' '.join(g.text for g in grp)
        runs = [r for g in grp for r in g.runs]
        if '→' in t or '←' in t:
            nch += 1
            if '→' in t and '←' in t: mixdir.append((i, t[:50]))
            # color of the text on each side of every arrow run must differ
            segs, cur = [], None
            for r in runs:
                if r.text.strip() in ('→', '←'):
                    segs.append(cur); cur = None
                elif r.text.strip() and cur is None:
                    cur = str(r.font.color.rgb) if r.font.color and r.font.color.type else None
            segs.append(cur)
            if any(a == b_ for a, b_ in zip(segs, segs[1:])): samecol.append((i, t[:50]))
        else:
            npl += 1
    if nch: chain_sl.append(i)
    elif npl: prose_sl.append(i)
    if nch and npl: mixed.append(i)
bad(f'bullets with both arrow directions: {mixdir[:4]}') if mixdir else ok('one arrow direction per bullet')
bad(f'same color on both sides of an arrow: {samecol[:4]}') if samecol else ok('colors differ across every arrow')
ok(f'{len(chain_sl)} arrow-chain slides')
bad(f'plain (non-chain) bullets on slides {sorted(set(prose_sl + mixed))[:10]}') if (prose_sl or mixed) else ok('no plain bullets: the deck is all arrow chains')

# ---------------------------------------------------------------- dividers --
print('SECTION DIVIDERS')
divs = [i for i, t in divider]
ok(f'dividers at slides {divs}') if divs else bad('no dividers')

# ------------------------------------------------------------------ frame --
print('FRAME SLIDES (v7: class title slide first, Thank you slide last)')
n_sl = len(p.slides)
if 1 in FRAME and n_sl in FRAME:
    t1 = ' '.join(sh.text_frame.text for sh in p.slides[0].shapes if sh.has_text_frame)
    tn = ' '.join(sh.text_frame.text for sh in p.slides[n_sl - 1].shapes if sh.has_text_frame)
    if '{{' in t1 + tn: bad('unfilled {{TOKEN}} on a frame slide')
    elif not re.search(r'Lecture \d{2}', t1): bad('title slide lacks "Lecture NN"')
    elif 'Thank you' not in tn: bad('last slide is not the Thank you slide')
    else: ok('title slide first, Thank you slide last, tokens filled: ' + re.search(r'Lecture \d{2}', t1).group(0))
else:
    bad(f'frame slides missing (title first: {1 in FRAME}, Thank you last: {n_sl in FRAME}): run frame.py and check Final.pptx')

print()
print(('ALL CHECKS PASSED' if not fail else f'{len(fail)} CHECK(S) FAILED'))
sys.exit(1 if fail else 0)
