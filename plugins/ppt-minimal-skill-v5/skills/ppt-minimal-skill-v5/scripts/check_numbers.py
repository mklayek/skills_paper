# -*- coding: utf-8 -*-
"""The user's rules: no equation numbers, page numbers, footers, figure numbers or section numbers.

Scans every text frame, table cell and native equation (OMML text) on every slide, plus the
slide XML for footer, date and slide-number placeholders or fields.
"""
import io, re, sys, zipfile
from pptx import Presentation

sys.stdout = io.open(1, "w", encoding="utf-8", closefd=False)
FN = sys.argv[1] if len(sys.argv) > 1 else "Deck.pptx"
PAT = {
    "figure number": re.compile(r"\bFig(?:ure|s|\.)?\s*\d", re.I),
    "table number": re.compile(r"\bTables?\s*\d", re.I),
    "section number": re.compile(r"\b(?:Section|Sec\.|Chapter)\s*\d|§\s*\d|\b7\.\d+(?:\.\d+)*\b(?!\s*(?:Hz|kHz|m|km|A|S|ms|μs|%|ppm|ft)\b)", re.I),
    "page number": re.compile(r"\b(?:p|pp|page)\.?\s*\d{2,}", re.I),
}
bad = []
prs = Presentation(FN)
for i, sl in enumerate(prs.slides, 1):
    texts = []
    for sh in sl.shapes:
        if sh.has_text_frame:
            texts += [p.text for p in sh.text_frame.paragraphs]
        if sh.has_table:   # a bare number in a data-table cell (e.g. a gain of 7.20) is a value, not a section number
            texts += [c.text for r in sh.table.rows for c in r.cells if not re.fullmatch(r"[−-]?\d+(\.\d+)?", c.text.strip())]
    # native equations: the text inside m:t elements, one string per paragraph
    for p in sl._element.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}p"):
        mt = [t.text or "" for t in p.iter("{http://schemas.openxmlformats.org/officeDocument/2006/math}t")]
        if mt:
            eq = "".join(mt); texts.append(eq)
            if re.search(r"\(\s*\d+(?:\.\d+)?[a-z]?\s*\)\s*$", eq):     # a number closing an equation line
                bad.append((i, "equation number", eq[:90]))
    for t in texts:
        for name, rx in PAT.items():
            if rx.search(t):
                bad.append((i, name, t[:90]))
z = zipfile.ZipFile(FN)
fields = [n for n in z.namelist() if n.startswith("ppt/slides/slide") and n.endswith(".xml")
          and re.search(rb'type="(?:sldNum|ftr|dt)"|<a:fld ', z.read(n))]
print(f"🔍 {len(prs.slides)} slides scanned")
print("✅ no footer, date or slide-number placeholders or fields" if not fields else f"❌ fields on {fields}")
if bad:
    for b in bad: print("❌", b)
else:
    print("✅ no equation, figure, table, section or page numbers anywhere")
