import json, sys
from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn
import omml
from eqcolor import color_equation
FN = sys.argv[1]; EQL = json.load(open('eqlist.json', encoding='utf-8')); p = Presentation(FN)
A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
MC = 'http://schemas.openxmlformats.org/markup-compatibility/2006'
A14 = 'http://schemas.microsoft.com/office/drawing/2010/main'
st = omml.Style('pptx', sz=2200, color='0B2A55'); done = 0
SRC = json.load(open('eqomml.json', encoding='utf-8'))
for i, sl in enumerate(p.slides, 1):
    info = EQL.get(str(i))
    if not info: continue
    body = [sh for sh in sl.shapes if sh.name == 'Body'][0]
    paras = [pa for k, pa in enumerate(body.text_frame.paragraphs)
             if k > 0 and pa.runs and pa.runs[0].font.name == 'Cambria Math']
    assert len(paras) == len(info['eqs']), (i, len(paras), len(info['eqs']))
    for pa, raw in zip(paras, info['eqs']):
        label, math, eqno = (None, None, None) if raw.startswith('§#') else omml.split_line(raw); p_ = pa._p
        runs = [r for r in p_ if r.tag in (qn('a:r'), qn('a:br'))]
        fbtxt = ''.join(''.join(r.itertext()) for r in runs if r.tag == qn('a:r')).strip()
        for r in runs: p_.remove(r)
        if label and fbtxt.startswith(label): fbtxt = fbtxt[len(label):].lstrip()
        fbline = ((label + '  ') if label else '') + ' '.join(fbtxt.split())
        xml = (f'<mc:AlternateContent xmlns:mc="{MC}" xmlns:a="{A}"><mc:Choice xmlns:a14="{A14}" Requires="a14"><a14:m>'
               + (SRC[raw[2:]] if raw.startswith('§#') else omml.to_omml(math, st, label, eqno, True).replace('<m:oMathPara xmlns:m=', f'<m:oMathPara xmlns:a="{A}" xmlns:m='))
               + f'</a14:m></mc:Choice><mc:Fallback><a:r><a:rPr lang="en-US" sz="2000"><a:solidFill><a:srgbClr val="0B2A55"/></a:solidFill>'
                 f'<a:latin typeface="Cambria Math"/></a:rPr><a:t>{omml.escape(fbline)}</a:t></a:r></mc:Fallback></mc:AlternateContent>')
        endp = p_.find(qn('a:endParaRPr')); el = etree.fromstring(xml)
        color_equation(el, done)              # LHS/RHS colors, alternating sets
        (endp.addprevious(el) if endp is not None else p_.append(el)); done += 1
p.save(FN); print('equations injected:', done)
