# -*- coding: utf-8 -*-
"""v7: put the class title slide first and the Thank you slide last.

    python frame.py Deck.pptx Final.pptx

Both slides are copied shape by shape from the skill's assets/frame_template.pptx (cut from the
user's class deck), so they keep every font, color, highlight and picture crop, and every text stays
an ordinary editable text box. Tokens are filled from deck_config.json:
    "method_name"  -> {{METHOD}}   e.g. "IP Method", "Electromagnetic Methods"
    "method_group" -> {{GROUP}}    e.g. "Electrical Methods"
    "lecture"      -> {{LECTURE}}  e.g. "07" (two digits)
Optional "title_photo", "title_book", "title_corner": image paths that replace the template pictures,
fitted inside the same box with the aspect ratio kept.
Deck.pptx (content only) is left unchanged, so the slide numbers used by the checks stay valid.
"""
import copy, io, json, os, sys
from pptx import Presentation
from pptx.oxml.ns import qn
from PIL import Image

sys.stdout = io.open(1, "w", encoding="utf-8", closefd=False)
SRC, OUT = sys.argv[1], sys.argv[2]
CFG = json.load(open("deck_config.json", encoding="utf-8"))
HERE = os.path.dirname(os.path.abspath(__file__))
TPL = CFG.get("frame_template") or next(
    (c for c in (os.path.join(HERE, "frame_template.pptx"), os.path.join(HERE, "..", "assets", "frame_template.pptx"))
     if os.path.exists(c)), None)
if not TPL:
    sys.exit("❌ frame_template.pptx not found: copy the skill's assets/frame_template.pptx into build/")
missing = [k for k in ("method_name", "method_group", "lecture") if not str(CFG.get(k, "")).strip()]
if missing:
    sys.exit(f"❌ deck_config.json lacks {missing}: ask the user for the method name, method group and lecture number")
TOK = {"{{METHOD}}": CFG["method_name"].strip(), "{{GROUP}}": CFG["method_group"].strip(),
       "{{LECTURE}}": str(CFG["lecture"]).strip().zfill(2)}
SWAP = {"TitlePhoto": CFG.get("title_photo"), "TitleBook": CFG.get("title_book"), "TitleCorner": CFG.get("title_corner")}

deck, tpl = Presentation(SRC), Presentation(TPL)
print(f"▶️  framing {SRC} ({len(deck.slides)} slides) with {os.path.basename(TPL)}")
print(f"🔧 title: {TOK['{{METHOD}}']} | {TOK['{{GROUP}}']} | Lecture {TOK['{{LECTURE}}']}")
layout = deck.slides[0].slide_layout


def fit_swap(pic_el, part, rid_old, path):
    """Point a picture at a new image, fitted inside the old box (aspect kept, centered), crop cleared."""
    _, rid = part.get_or_add_image_part(path)
    pic_el.find('.//' + qn('a:blip')).set(qn('r:embed'), rid)
    src = pic_el.find('.//' + qn('a:srcRect'))
    if src is not None:
        for k in list(src.attrib): del src.attrib[k]
    off, ext = pic_el.find('.//' + qn('a:off')), pic_el.find('.//' + qn('a:ext'))
    x, y, w, h = (int(off.get('x')), int(off.get('y')), int(ext.get('cx')), int(ext.get('cy')))
    iw, ih = Image.open(path).size
    s = min(w / iw, h / ih); nw, nh = int(iw * s), int(ih * s)
    off.set('x', str(x + (w - nw) // 2)); off.set('y', str(y + (h - nh) // 2))
    ext.set('cx', str(nw)); ext.set('cy', str(nh))


def clone(src_slide, index):
    """Copy every shape of a template slide into a new deck slide placed at `index`."""
    s = deck.slides.add_slide(layout)
    for ph in list(s.placeholders): ph._element.getparent().remove(ph._element)
    tree = s.shapes._spTree
    for el in src_slide.shapes._spTree.iterchildren():
        if el.tag in (qn('p:nvGrpSpPr'), qn('p:grpSpPr')): continue
        new = copy.deepcopy(el); tree.append(new)
        for blip in new.iter(qn('a:blip')):
            rid = blip.get(qn('r:embed'))
            img = src_slide.part.related_part(rid)
            _, nrid = s.part.get_or_add_image_part(io.BytesIO(img.blob))
            blip.set(qn('r:embed'), nrid)
        name = new.find('.//' + qn('p:cNvPr')).get('name')
        if SWAP.get(name):
            fit_swap(new, s.part, None, SWAP[name]); print(f"🖼️  {name} replaced by {SWAP[name]}")
        for t in new.iter(qn('a:t')):
            if t.text and '{{METHOD}}' in t.text:
                # a long method name ("Electromagnetic Methods") must stay on one line in the 5.9 in box
                rpr = t.getparent().find(qn('a:rPr'))
                sz = min(5400, int(5.7 * 72 / (len(TOK['{{METHOD}}']) * 0.52)) * 100)
                if rpr is not None and sz < int(rpr.get('sz', '5400')): rpr.set('sz', str(sz))
            for k, v in TOK.items():
                if t.text and k in t.text: t.text = t.text.replace(k, v)
    ids = deck.slides._sldIdLst; sid = ids[-1]; ids.remove(sid); ids.insert(index, sid)
    return s


clone(tpl.slides[0], 0)
clone(tpl.slides[1], len(deck.slides._sldIdLst))
left = [sh.name for sl in (deck.slides[0], deck.slides[-1]) for sh in sl.shapes
        if sh.has_text_frame and '{{' in sh.text_frame.text]
if left: sys.exit(f"❌ unfilled tokens in {left}")
deck.save(OUT)
print(f"✅ {OUT} written: title slide + {len(deck.slides) - 2} content slides + Thank you slide")
