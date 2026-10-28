import sys, zipfile, shutil
from pptx import Presentation
from pptx.oxml.ns import qn
from lxml import etree
FN = sys.argv[1]; p = Presentation(FN)


def kind(para, k):
    r = para.runs[0]
    if k == 0 and (r.font.underline or (r.font.size and r.font.size.pt == 18)): return 'sub'
    if r.font.name == 'Cambria Math': return 'eq'
    if r.font.size and r.font.size.pt == 18 and r.font.name == 'Times New Roman': return 'sym'   # v7 symbol line
    if r.font.name in ('Cambria', 'Arial') and r.font.size and r.font.size.pt == 20: return 'sym'
    return 'bul'


for sl in p.slides:
    for sh in sl.shapes:
        if not sh.has_text_frame: continue
        for para in sh.text_frame.paragraphs:
            pps = para._p.findall(qn('a:pPr'))
            for e in pps[1:]: para._p.remove(e)
            if pps and para._p.index(pps[0]) != 0:
                para._p.remove(pps[0]); para._p.insert(0, pps[0])
        if sh.name != 'Body': continue
        for k, para in enumerate(sh.text_frame.paragraphs):
            if not para.runs: continue
            t = kind(para, k); pPr = para._p.get_or_add_pPr()
            for tag in ('a:buNone', 'a:buChar', 'a:buFont', 'a:buAutoNum', 'a:spcAft',
                        'a:spcBef', 'a:buClr', 'a:buSzPct'):
                for e in pPr.findall(qn(tag)): pPr.remove(e)

            def spc(tag, v):
                e = etree.SubElement(pPr, qn(tag))
                etree.SubElement(e, qn('a:spcPts')).set('val', str(v))

            if t == 'sub':
                pPr.set('marL', '0'); pPr.set('indent', '0'); spc('a:spcAft', 800)
                etree.SubElement(pPr, qn('a:buNone'))
            elif t == 'eq':
                pPr.set('marL', '457200'); pPr.set('indent', '0')
                spc('a:spcBef', 200); spc('a:spcAft', 400)
                etree.SubElement(pPr, qn('a:buNone'))
            elif t == 'sym':
                # v7: no bullet, aligned with the equation it explains
                pPr.set('marL', '457200'); pPr.set('indent', '0'); spc('a:spcAft', 400)
                etree.SubElement(pPr, qn('a:buNone'))
            else:
                pPr.set('marL', '285750'); pPr.set('indent', '-285750'); spc('a:spcAft', 600)
                etree.SubElement(pPr, qn('a:buFont')).set('typeface', 'Arial')
                etree.SubElement(pPr, qn('a:buChar')).set('char', '•')
            ln = pPr.find(qn('a:lnSpc'))
            if ln is not None: pPr.remove(ln); pPr.insert(0, ln)
p.save(FN)
zin = zipfile.ZipFile(FN); zout = zipfile.ZipFile('tmp.pptx', 'w', zipfile.ZIP_DEFLATED)
for it in zin.infolist():
    d = zin.read(it.filename)
    if it.filename == '[Content_Types].xml':
        t = d.decode()
        if 'Extension="jpg"' not in t:
            t = t.replace('<Default ', '<Default Extension="jpg" ContentType="image/jpeg"/><Default ', 1)
        d = t.encode()
    zout.writestr(it, d)
zin.close(); zout.close(); shutil.move('tmp.pptx', FN)
print('post-processing done')
