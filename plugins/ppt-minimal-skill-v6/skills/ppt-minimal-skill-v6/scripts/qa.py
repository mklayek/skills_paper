# -*- coding: utf-8 -*-
"""Checks the built deck against the house rules and this deck's own two extra rules:
no equation numbers and no slide numbers."""
import re, sys, zipfile
from pptx import Presentation
from pptx.util import Pt

FN = sys.argv[1] if len(sys.argv) > 1 else 'Deck.pptx'
p = Presentation(FN)
fail = []
ok = lambda m: print('  OK   ' + m)


def bad(m):
    fail.append(m); print('  FAIL ' + m)


# ---------------------------------------------------------------- structure --
print('STRUCTURE')
boxes_bad = []
for i, sl in enumerate(p.slides, 1):
    names = [sh.name for sh in sl.shapes if sh.has_text_frame]
    tx = [sh for sh in sl.shapes if sh.has_text_frame]
    if names.count('Title') != 1: boxes_bad.append(f'slide {i}: {names.count("Title")} Title')
    if names.count('Body') != 1: boxes_bad.append(f'slide {i}: {names.count("Body")} Body')
    extra = [sh for sh in tx if sh.name not in ('Title', 'Body')]
    # an image caption is the only other text box allowed
    for e in extra:
        t = e.text_frame.text.strip()
        if not t:          # the image border rectangle carries an empty text frame
            continue
        runs = [r for para in e.text_frame.paragraphs for r in para.runs]
        cap = bool(runs) and all(r.font.italic and r.font.size and r.font.size.pt == 12
                                 for r in runs)
        if not (cap or t.startswith('Section ')):
            boxes_bad.append(f'slide {i}: unexpected text box {t[:40]!r}')
print(f'  slides: {len(p.slides)}')
bad('; '.join(boxes_bad)) if boxes_bad else ok('every slide has exactly one Title and one Body, no stray boxes')

# ------------------------------------------------------------------ titles --
print('TITLES')
titles = []
for i, sl in enumerate(p.slides, 1):
    t = [sh for sh in sl.shapes if sh.name == 'Title']
    titles.append((i, t[0].text_frame.text.strip() if t else ''))
content = [(i, t) for i, t in titles if t and t[0] in '◆■●']
divider = [(i, t) for i, t in titles if t and t[0] not in '◆■●']
numbered = [(i, t) for i, t in content if re.match(r'^[◆■●]\s+\d', t)]
secnum = [(i, t) for i, t in content if re.search(r'\b\d+\.\d+\b', t)]
ok(f'{len(content)} content titles carry a level marker, {len(divider)} dividers')
bad(f'titles with a leading number: {numbered}') if numbered else ok('no numbered titles')
bad(f'titles with a source section number: {secnum}') if secnum else ok('no source section numbers in titles')

# ------------------------------------------- slide numbers and footers (none) --
print('NO SLIDE NUMBERS / FOOTERS')
with zipfile.ZipFile(FN) as z:
    sld = [n for n in z.namelist() if n.startswith('ppt/slides/slide')]
    fld = [n for n in sld if b'<a:fld' in z.read(n) and b'type="slidenum"' in z.read(n)]
bad(f'slide-number fields in {fld}') if fld else ok('no slide-number field anywhere')

# ------------------------------------------------- equation numbers (none) --
print('NO EQUATION NUMBERS')
eqno = []
for i, sl in enumerate(p.slides, 1):
    for sh in sl.shapes:
        if not sh.has_text_frame: continue
        for para in sh.text_frame.paragraphs:
            txt = para.text
            if re.search(r'\(\s*\d+[a-e]?\s*\)\s*$', txt) and not re.search(r'(Fig|Table)', txt):
                eqno.append((i, txt[-40:]))
bad(f'trailing equation numbers: {eqno[:4]}') if eqno else ok('no equation numbers on any line')
chap = [(i, pa.text) for i, sl in enumerate(p.slides, 1) for sh in sl.shapes
        if sh.has_text_frame for pa in sh.text_frame.paragraphs
        if re.search(r'\(\d+\.\d+\)', pa.text)]
bad(f'chapter-prefixed numbers: {chap[:3]}') if chap else ok('no chapter-prefixed equation numbers')

# -------------------------------------------------------------- raw markup --
print('NO RAW MARKUP')
raw = [(i, pa.text[:50]) for i, sl in enumerate(p.slides, 1) for sh in sl.shapes
       if sh.has_text_frame for pa in sh.text_frame.paragraphs
       if '_{' in pa.text or '^{' in pa.text or '**' in pa.text or re.search(r'[A-Za-z]_[A-Za-zρθφ]', pa.text)]
bad(f'raw markup left: {raw[:4]}') if raw else ok('no _{ ^{ or ** left in the text')

# --------------------------------------------------------------- US English --
print('US ENGLISH')
WHITE = set('noise source four hour contour precise rise expertise promise filled controlled '
            'called drilling drilled flour pour your our tour raised praised noised poised '
            'revised advised devised comprised supervised'.split())
pat = re.compile(r'\b\w+(?:isation|ised|ising|yse|ysed|olour|aviour|metre|entre|entres|'
                 r'ramme|ibre|itre)\b|\bgrey\b|\bwhilst\b|\bamongst\b|\balum inium\b', re.I)
hits = []
for i, sl in enumerate(p.slides, 1):
    for sh in sl.shapes:
        if sh.has_text_frame:
            for m in pat.finditer(sh.text_frame.text):
                if m.group(0).lower() not in WHITE: hits.append((i, m.group(0)))
        if sh.has_table:
            for r in sh.table.rows:
                for c in r.cells:
                    for m in pat.finditer(c.text):
                        if m.group(0).lower() not in WHITE: hits.append((i, m.group(0)))
bad(f'British spellings: {sorted(set(h[1] for h in hits))}') if hits else ok('US English throughout')

# --------------------------------------------------------------- subtitles --
print('SUBTITLES (v2: only where they add information)')
sub = lambda pa: bool(pa.runs) and bool(pa.runs[0].font.underline)
nsub = [i for i, sl in enumerate(p.slides, 1) for sh in sl.shapes
        if sh.name == 'Body' and sh.text_frame.paragraphs and sub(sh.text_frame.paragraphs[0])]
ok(f'{len(nsub)} slides keep a subtitle (title slide and named case histories)')

# ------------------------------------------------------------ title fonts --
print('TITLE FONTS (v2: Copperplate default, one font and color per section)')
sec, seen, tbad = 0, {}, []
for i, sl in enumerate(p.slides, 1):
    tt = [sh for sh in sl.shapes if sh.name == 'Title'][0]
    txt = tt.text_frame.text.strip()
    runs = tt.text_frame.paragraphs[0].runs
    r = runs[-1]
    sig = (r.font.name, str(r.font.color.rgb))
    if txt[:1] not in '◆■●':
        sec += 1
    if txt.endswith('Summary'):
        key = 'summary'
    else:
        key = sec
    if key in seen and seen[key] != sig: tbad.append((i, sig, seen[key]))
    seen.setdefault(key, sig)
fonts = {k: v[0] for k, v in seen.items()}
if seen.get(0, ('',))[0] != 'Copperplate Gothic Bold' or seen.get('summary', ('',))[0] != 'Copperplate Gothic Bold':
    tbad.append(('default', seen.get(0), seen.get('summary')))
distinct = len(set(v for k, v in seen.items() if k not in (0, 1, 'summary')))
bad(f'title font breaks: {tbad[:4]}') if tbad else ok(f'title fonts per section: {fonts}')
ok(f'{distinct} distinct font/color pairs across sections 2 to {sec}')

# ------------------------------------------------- semicolons and em dashes --
print('WRITING (no semicolons, no em dashes)')
pun = []
for i, sl in enumerate(p.slides, 1):
    for sh in sl.shapes:
        texts = [sh.text_frame.text] if sh.has_text_frame else []
        if sh.has_table: texts += [c.text for r in sh.table.rows for c in r.cells]
        for x in texts:
            if ';' in x or '—' in x: pun.append((i, x[:60]))
bad(f'semicolons or em dashes: {pun[:4]}') if pun else ok('no semicolons and no em dashes in slide text')

# ------------------------------------------------------------ bullet fonts --
print('BULLET FONT CYCLE')
FONTS = ['Times New Roman', 'Comic Sans MS', 'Bookman Old Style']
wrong = []
for i, sl in enumerate(p.slides, 1):
    b = [sh for sh in sl.shapes if sh.name == 'Body']
    if not b: continue
    if titles[i - 1][1][:1] not in '◆■●': continue      # divider slides carry no bullets
    k = 0
    for j, pa in enumerate(b[0].text_frame.paragraphs):
        if not pa.runs or (j == 0 and sub(pa)): continue
        f = pa.runs[0].font
        if f.name in ('Cambria Math', 'Cambria', 'Arial'): continue
        if f.name != FONTS[k % 3]: wrong.append((i, j, f.name, FONTS[k % 3]))
        if f.size and f.size != Pt(22): wrong.append((i, j, 'size', f.size.pt))
        k += 1
bad(f'bullet font/size breaks: {wrong[:4]}') if wrong else ok('bullets follow the 3-font cycle at 22 pt')

# ------------------------------------------------------------ arrow chains --
print('ARROW CHAINS (v6)')
chain_sl, prose_sl, mixdir, samecol, mixed = [], [], [], [], []
for i, sl in enumerate(p.slides, 1):
    if titles[i - 1][1][:1] not in '◆■●': continue
    b = [sh for sh in sl.shapes if sh.name == 'Body']
    if not b: continue
    nch = npl = 0
    for j, pa in enumerate(b[0].text_frame.paragraphs):
        if not pa.runs or (j == 0 and sub(pa)): continue
        if pa.runs[0].font.name in ('Cambria Math', 'Cambria', 'Arial'): continue
        t = pa.text
        if '→' in t or '←' in t:
            nch += 1
            if '→' in t and '←' in t: mixdir.append((i, t[:50]))
            # color of the text on each side of every arrow run must differ
            segs, cur = [], None
            for r in pa.runs:
                if r.text.strip() in ('→', '←'):
                    segs.append(cur); cur = None
                elif r.text.strip() and cur is None:
                    cur = str(r.font.color.rgb) if r.font.color and r.font.color.type else None
            segs.append(cur)
            if any(a == b_ for a, b_ in zip(segs, segs[1:])): samecol.append((i, t[:50]))
        else:
            npl += 1
    if nch: chain_sl.append(i)
    elif npl: prose_sl.append(i)
    if nch and npl: mixed.append(i)
bad(f'bullets with both arrow directions: {mixdir[:4]}') if mixdir else ok('one arrow direction per bullet')
bad(f'same color on both sides of an arrow: {samecol[:4]}') if samecol else ok('colors differ across every arrow')
ok(f'{len(chain_sl)} arrow-chain slides, {len(prose_sl)} source-style slides')
if mixed: print(f'  NOTE slides mixing chain and plain bullets: {mixed[:8]} (fine for a lead-in sentence)')
if chain_sl and not prose_sl or prose_sl and not chain_sl:
    print('  NOTE the deck is all one style: v6 expects a mix unless the user asked otherwise')

# ---------------------------------------------------------------- dividers --
print('SECTION DIVIDERS')
divs = [i for i, t in divider]
ok(f'dividers at slides {divs}') if divs else bad('no dividers')

print()
print(('ALL CHECKS PASSED' if not fail else f'{len(fail)} CHECK(S) FAILED'))
sys.exit(1 if fail else 0)
