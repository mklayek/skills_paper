"""Render Deck_linear.pptx and report slide count, (continued) slides, thin ones and any overflow."""
import io, subprocess, sys
import pymupdf as fitz
from pptx import Presentation
sys.stdout = io.open(1, "w", encoding="utf-8", closefd=False)
subprocess.run([r"C:\Program Files\LibreOffice\program\soffice.exe", "--headless", "--convert-to", "pdf",
                "--outdir", "../extract", "Deck_linear.pptx"], check=True, capture_output=True)
prs = Presentation("Deck_linear.pptx"); doc = fitz.open("../extract/Deck_linear.pdf")
cont = thin = 0; over = []
for i, (sl, pg) in enumerate(zip(prs.slides, doc), 1):
    t = [sh for sh in sl.shapes if sh.name == "Title"][0].text_frame.text
    body = [sh for sh in sl.shapes if sh.name == "Body"][0].text_frame
    half = any(sh.shape_type == 13 or sh.has_table for sh in sl.shapes)
    blocks = [b for b in pg.get_text("blocks") if b[0] < 72 * (6.6 if half else 12.8) and b[1] > 72 * 1.2]
    low = max((b[3] for b in blocks), default=0) / 72
    if "(continued)" in t:
        cont += 1
        if len(body.paragraphs) - 1 <= 2: thin += 1; print("thin", i, len(body.paragraphs) - 1, t[:50])
    if low > 6.95: over.append((i, round(low, 2), t[:50]))
print(f"slides {len(prs.slides)}, cont {cont}, thin cont {thin}, overflow (linear render) {len(over)}")
for o in over: print("   over", o)
