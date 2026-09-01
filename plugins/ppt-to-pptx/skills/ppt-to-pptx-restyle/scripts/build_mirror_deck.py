# -*- coding: utf-8 -*-
"""Build a slide-for-slide ppt-mkl .deck mirror from a source .pptx.
Usage:  python build_mirror_deck.py <src.pptx> <out_dir> <overrides.txt> <deck.out>

For each source slide (2..N; slide 1 = title -> ppt-mkl flags):
  * If an override block "=== SLIDE n ===" exists (from agent transcription of screenshot
    text/equations), use its TITLE / bullets / $$equations$$ and its KEEP: figure list.
  * Else use the auto-extracted verbatim text (from the pptx) and montage ALL of the slide's
    extracted images.
Multiple KEEP figures (or multiple auto images) are montaged into one image (ppt-mkl = 1 fig/slide).
Slides whose title is "thank you" are skipped (ppt-mkl auto-appends a Thank-You).

overrides.txt block format (one per slide):
  === SLIDE n ===
  TITLE: <title>
  - <verbatim bullet>
  $$ <LaTeX> $$
  KEEP: <fig1> , <fig2> , ...        # optional; genuine diagrams only (omit text-box screenshots)
"""
import os, sys, re, math, glob
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from PIL import Image

SRC, OUT, OVR, DECK = sys.argv[1], sys.argv[2], sys.argv[3], sys.argv[4]
FIG = os.path.join(OUT, "figures")
MON = os.path.join(FIG, "montage"); os.makedirs(MON, exist_ok=True)

CHROME = {"mkl"}
def is_chrome(t):
    s = t.strip()
    return s == "" or s.isdigit() or s.lower() in CHROME

def auto_text(slide):
    title = ""
    try:
        if slide.shapes.title and slide.shapes.title.text.strip():
            title = slide.shapes.title.text.strip()
    except Exception: pass
    out = []
    def walk(shapes):
        for sh in shapes:
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP: walk(sh.shapes); continue
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    out.append((p.level or 0, ("".join(r.text for r in p.runs) or p.text)))
    walk(slide.shapes)
    if not title:
        for _, t in out:
            if t.strip() and not t.strip().isdigit() and t.strip().lower() != "mkl":
                title = t.strip(); break
    return title, out

def parse_overrides(path):
    if not path or not os.path.exists(path): return {}
    blocks = {}
    cur = None
    for raw in open(path, encoding="utf-8"):
        line = raw.rstrip("\n")
        m = re.match(r"=== SLIDE (\d+) ===", line.strip())
        if m:
            cur = int(m.group(1)); blocks[cur] = {"title": "", "body": [], "keep": []}
            continue
        if cur is None: continue
        if line.startswith("TITLE:"):
            blocks[cur]["title"] = line[6:].strip()
        elif line.startswith("KEEP:"):
            blocks[cur]["keep"] = [x.strip() for x in line[5:].split(",") if x.strip()]
        elif line.startswith("[DIAGRAM"):
            continue
        elif line.strip():
            blocks[cur]["body"].append(line.rstrip())
    return blocks

def slide_imgs(i):
    return sorted(f for f in glob.glob(os.path.join(FIG, f"slide{i:02d}_r*.*")))

def resolve_keep(keep):
    out = []
    for k in keep:
        if os.path.isabs(k) and os.path.exists(k): out.append(k); continue
        cand = os.path.join(FIG, os.path.basename(k))
        if os.path.exists(cand): out.append(cand); continue
        g = glob.glob(os.path.join(FIG, os.path.basename(k).split(".")[0] + ".*"))
        if g: out.append(g[0])
    return out

def montage(paths, out):
    ims = [Image.open(p).convert("RGB") for p in paths]
    n = len(ims); cols = min(n, 3); rows = math.ceil(n/cols); H = 320
    ims = [im.resize((max(1, int(im.width*H/im.height)), H)) for im in ims]
    colw = max(im.width for im in ims); gap = 16
    sh = Image.new("RGB", (cols*colw+gap*(cols+1), rows*H+gap*(rows+1)), "white")
    for k, im in enumerate(ims):
        r, c = divmod(k, cols)
        sh.paste(im, (gap+c*(colw+gap)+(colw-im.width)//2, gap+r*(H+gap)))
    sh.save(out); return out

def fig_line(i, files):
    files = [f for f in files if os.path.exists(f)]
    if not files: return None
    if len(files) == 1:
        p = files[0].replace("\\", "/"); return f"!fig: {p} 55%"
    m = montage(files, os.path.join(MON, f"slide{i:02d}_montage.png")).replace("\\", "/")
    return f"!fig: {m} 85%"

ovr = parse_overrides(OVR)
prs = Presentation(SRC)
out = []
for i, slide in enumerate(prs.slides, 1):
    if i == 1: continue
    if i in ovr:
        b = ovr[i]; title = b["title"] or "Slide"
        if title.strip().lower() == "thank you": continue
        out.append(f"\n== {title}")
        out += b["body"]
        fl = fig_line(i, resolve_keep(b["keep"]) if b["keep"] else [])
        if fl: out.append(fl)
    else:
        title, body = auto_text(slide)
        if title.strip().lower() == "thank you": continue
        if not title.strip(): title = "Slide"
        out.append(f"\n== {title}")
        for lvl, t in body:
            if is_chrome(t) or t.strip().lower() == title.strip().lower(): continue
            out.append(("-" if lvl == 0 else "--" if lvl == 1 else "---") + " " + t.strip())
        fl = fig_line(i, slide_imgs(i))
        if fl: out.append(fl)

os.makedirs(os.path.dirname(DECK) or ".", exist_ok=True)
open(DECK, "w", encoding="utf-8").write("\n".join(out).lstrip() + "\n")
print("wrote", DECK, "| slides:", sum(1 for l in out if l.startswith("== ")))
